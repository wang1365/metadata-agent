import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
import xlrd
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils.cell import range_boundaries

MAX_CELLS = 1_000_000


def preview_region(path, sheet_name, region):
    """Read bounded worksheet presentation data; no layout inference or formula execution."""
    left, top, right, bottom = range_boundaries(region)
    if not all([left, top, right, bottom]) or min(left, top) < 1 or right < left or bottom < top:
        raise ValueError("INVALID_REGION")
    if (right - left + 1) * (bottom - top + 1) > 2000:
        raise ValueError("REGION_LIMIT: use at most 2000 cells")
    book = workbook(path)
    try:
        xlsx = isinstance(book, openpyxl.Workbook)
        sheet = book[sheet_name] if xlsx else book.sheet_by_name(sheet_name)
        max_row, max_col = (sheet.max_row, sheet.max_column) if xlsx else (sheet.nrows, sheet.ncols)
        if right > max_col or bottom > max_row:
            raise ValueError("REGION_OUT_OF_RANGE")
        theme_colors = {}
        if xlsx and book.loaded_theme:
            ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
            try:
                theme = ElementTree.fromstring(book.loaded_theme)
                scheme = theme.find(f"{ns}themeElements/{ns}clrScheme")
                for index, key in enumerate(
                    (
                        "lt1",
                        "dk1",
                        "lt2",
                        "dk2",
                        "accent1",
                        "accent2",
                        "accent3",
                        "accent4",
                        "accent5",
                        "accent6",
                        "hlink",
                        "folHlink",
                    )
                ):
                    element = scheme.find(ns + key) if scheme is not None else None
                    if element is not None and len(element):
                        theme_colors[index] = element[0].get("lastClr") or element[0].get("val")
            except ElementTree.ParseError:
                pass

        def color(value):
            if value is None:
                return None
            rgb = value.rgb if value.type == "rgb" else None
            if value.type == "indexed" and 0 <= value.indexed < len(COLOR_INDEX):
                rgb = COLOR_INDEX[value.indexed]
            if value.type == "theme":
                rgb = theme_colors.get(value.theme)
            if not isinstance(rgb, str) or len(rgb) not in (6, 8):
                return None
            rgb = rgb[-6:]
            try:
                channels = [int(rgb[i : i + 2], 16) for i in (0, 2, 4)]
            except ValueError:
                return None
            tint = value.tint or 0
            return "#" + "".join(f"{round(c * (1 + tint) if tint < 0 else c + (255 - c) * tint):02x}" for c in channels)

        def display(value, fmt=""):
            if value is None:
                return ""
            if isinstance(value, (date, datetime)):
                return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
            if isinstance(value, bool):
                return "TRUE" if value else "FALSE"
            if isinstance(value, (int, float)):
                simple = fmt.split(";")[0]
                if len(simple) <= 100 and re.fullmatch(r"0+(\.0+)?%?", simple):
                    decimals = len(simple.split(".")[1].rstrip("%")) if "." in simple else 0
                    if simple.endswith("%"):
                        return f"{value * 100:.{decimals}f}%"
                    if "." in simple:
                        return f"{value:.{decimals}f}"
                    return str(int(value)).zfill(len(simple))
            return str(value)[:4000]

        columns, rows, cells = [], [], []
        for c in range(left, right + 1):
            letter = openpyxl.utils.get_column_letter(c)
            if xlsx:
                dimension = next(
                    (d for d in sheet.column_dimensions.values() if (d.min or 0) <= c <= (d.max or 0)), None
                )
                width = (dimension.width if dimension else sheet.sheet_format.defaultColWidth or 13) * 7 + 5
                hidden = bool(dimension and dimension.hidden)
            else:
                dimension = sheet.colinfo_map.get(c - 1)
                width = (dimension.width / 256 if dimension else 13) * 7 + 5
                hidden = bool(dimension and dimension.hidden)
            columns.append({"index": c, "label": letter, "width": max(20, min(1000, width)), "hidden": hidden})
        for r in range(top, bottom + 1):
            dimension = sheet.row_dimensions.get(r) if xlsx else sheet.rowinfo_map.get(r - 1)
            height = (dimension.height or 15) if xlsx and dimension else (dimension.height / 20 if dimension else 15)
            rows.append(
                {
                    "index": r,
                    "height": max(20, min(500, height * 4 / 3)),
                    "hidden": bool(dimension and dimension.hidden),
                }
            )
        merges = (
            [range_boundaries(str(m)) for m in sheet.merged_cells.ranges]
            if xlsx
            else [(c + 1, r + 1, ce, re) for r, re, c, ce in sheet.merged_cells]
        )
        merged = {}
        covered = set()
        for ml, mt, mr, mb in merges:
            if ml > right or mr < left or mt > bottom or mb < top:
                continue
            cl, ct, cr, cb = max(left, ml), max(top, mt), min(right, mr), min(bottom, mb)
            merged[(ct, cl)] = (ml, mt, cb - ct + 1, cr - cl + 1)
            covered.update((r, c) for r in range(ct, cb + 1) for c in range(cl, cr + 1) if (r, c) != (ct, cl))
        for r in range(top, bottom + 1):
            for c in range(left, right + 1):
                if (r, c) in covered:
                    continue
                ac, ar, rowspan, colspan = merged.get((r, c), (c, r, 1, 1))
                style = {}
                if xlsx:
                    cell = sheet.cell(ar, ac)
                    font, alignment = cell.font, cell.alignment
                    value, fmt = cell.value, cell.number_format
                    style = {
                        "fontFamily": font.name,
                        "fontSize": f"{font.sz or 11}pt",
                        "fontWeight": "bold" if font.bold else "normal",
                        "fontStyle": "italic" if font.italic else "normal",
                        "textDecoration": "underline" if font.underline else "none",
                        "color": color(font.color),
                        "backgroundColor": color(cell.fill.fgColor) if cell.fill.patternType == "solid" else None,
                        "textAlign": alignment.horizontal
                        if alignment.horizontal in ("left", "center", "right", "justify")
                        else None,
                        "verticalAlign": {"center": "middle", "top": "top", "bottom": "bottom"}.get(alignment.vertical),
                        "whiteSpace": "pre-wrap" if alignment.wrap_text else "pre",
                    }
                    for edge in ("top", "bottom", "left", "right"):
                        side = getattr(cell.border, edge)
                        if side and side.style:
                            line = (
                                "dashed"
                                if "dash" in side.style.lower()
                                else "dotted"
                                if side.style == "dotted"
                                else "solid"
                            )
                            style["border" + edge.capitalize()] = (
                                f"{2 if side.style in ('medium', 'thick', 'double') else 1}px {line} {color(side.color) or '#64748b'}"
                            )
                else:
                    cell = sheet.cell(ar - 1, ac - 1)
                    value = cell.value
                    xf = book.xf_list[cell.xf_index]
                    font = book.font_list[xf.font_index]
                    fmt = book.format_map[xf.format_key].format_str

                    def xls_color(index):
                        rgb = book.colour_map.get(index)
                        return "#" + "".join(f"{part:02x}" for part in rgb) if rgb else None

                    style = {
                        "fontFamily": font.name,
                        "fontSize": f"{font.height / 20}pt",
                        "fontWeight": "bold" if font.bold else "normal",
                        "fontStyle": "italic" if font.italic else "normal",
                        "color": xls_color(font.colour_index),
                        "backgroundColor": xls_color(xf.background.pattern_colour_index)
                        if xf.background.fill_pattern == 1
                        else None,
                        "textAlign": {1: "left", 2: "center", 3: "right"}.get(xf.alignment.hor_align),
                        "whiteSpace": "pre-wrap" if xf.alignment.text_wrapped else "pre",
                    }
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                cells.append(
                    {
                        "row": r,
                        "column": c,
                        "value": display(value, fmt),
                        "rowSpan": rowspan,
                        "colSpan": colspan,
                        "style": {k: v for k, v in style.items() if v is not None},
                    }
                )
        return {
            "sheet": sheet_name,
            "region": region,
            "columns": columns,
            "rows": rows,
            "cells": cells,
            "warnings": ["预览保留行列、合并与常见样式；不执行公式，暂不还原图片、图表、条件格式或复杂数字格式。"],
        }
    finally:
        book.close() if isinstance(book, openpyxl.Workbook) else book.release_resources()


def workbook(path):
    path = Path(path)
    if path.suffix == ".xlsx":
        if not zipfile.is_zipfile(path):
            raise ValueError("INVALID_XLSX")
        with zipfile.ZipFile(path) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 200 * 1024 * 1024:
                raise ValueError("WORKBOOK_EXPANSION_LIMIT")
            if len(archive.infolist()) > 10000:
                raise ValueError("WORKBOOK_ENTRY_LIMIT")
        book = openpyxl.load_workbook(path, read_only=False, data_only=False, keep_links=False)
        try:
            for sheet in book:
                if sheet.max_row * sheet.max_column > MAX_CELLS:
                    raise ValueError("WORKBOOK_CELL_LIMIT")
            if sum(s.max_row * s.max_column for s in book) > MAX_CELLS:
                raise ValueError("WORKBOOK_CELL_LIMIT")
            return book
        except BaseException:
            book.close()
            raise
    if path.suffix != ".xls":
        raise ValueError("UNSUPPORTED_FILE_TYPE")
    book = xlrd.open_workbook(path, formatting_info=True)
    if sum(s.nrows * s.ncols for s in book.sheets()) > MAX_CELLS:
        book.release_resources()
        raise ValueError("WORKBOOK_CELL_LIMIT")
    return book


def inspect_workbook(path):
    book = workbook(path)
    try:
        if isinstance(book, openpyxl.Workbook):
            return [
                {
                    "name": s.title,
                    "rows": s.max_row,
                    "columns": s.max_column,
                    "merged": [str(r) for r in s.merged_cells.ranges],
                }
                for s in book
            ]
        return [
            {"name": s.name, "rows": s.nrows, "columns": s.ncols, "merged": [list(r) for r in s.merged_cells]}
            for s in book.sheets()
        ]
    finally:
        book.close() if isinstance(book, openpyxl.Workbook) else book.release_resources()


def read_region(path, sheet_name, region):
    left, top, right, bottom = range_boundaries(region)
    if (
        not all([left, top, right, bottom])
        or min(left, top) < 1
        or right < left
        or bottom < top
        or (right - left + 1) * (bottom - top + 1) > 2000
    ):
        raise ValueError("REGION_LIMIT: use at most 2000 cells")
    book = workbook(path)
    try:
        if isinstance(book, openpyxl.Workbook):
            sheet = book[sheet_name]
            if right > sheet.max_column or bottom > sheet.max_row:
                raise ValueError("REGION_OUT_OF_RANGE")
            return [
                {"cell": cell.coordinate, "value": str(cell.value)[:4000]}
                for row in sheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right)
                for cell in row
                if cell.value is not None
            ]
        sheet = book.sheet_by_name(sheet_name)
        if right > sheet.ncols or bottom > sheet.nrows:
            raise ValueError("REGION_OUT_OF_RANGE")
        return [
            {"row": r + 1, "column": c + 1, "value": str(sheet.cell_value(r, c))[:4000]}
            for r in range(top - 1, bottom)
            for c in range(left - 1, right)
            if sheet.cell_value(r, c) != ""
        ]
    finally:
        book.close() if isinstance(book, openpyxl.Workbook) else book.release_resources()

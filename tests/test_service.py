import io

import openpyxl
import pytest

from dsh.plugins.metadata_tools.capabilities import Capabilities
from metadata_service.models import Findings


def upload_workbook(client):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "字典"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "用户表"
    sheet.append(["字段", "类型", "说明"])
    sheet.append(["id", "bigint", "主键"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = client.post("/api/v1/sources/excel", files={"file": ("dictionary.xlsx", buffer.getvalue())})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def prepared(client, store):
    source_id = upload_workbook(client)
    job_id = client.post("/api/v1/analyses", json={"source_id": source_id}).json()["task_id"]
    store.claim(job_id)
    tools = Capabilities(store, job_id)
    evidence_id = tools.read_region("字典", "A1:C3")["evidence_ids"][0]
    findings = {
        "tables": [
            {
                "name": "users",
                "columns": [{"name": "id", "type": "BIGINT", "evidence_ids": [evidence_id]}],
                "keys": [{"kind": "primary", "columns": ["id"]}],
                "evidence_ids": [evidence_id],
            }
        ]
    }
    return job_id, tools, findings


def test_excel_confirmation_publish_and_qa_scope(client, store):
    job_id, tools, findings = prepared(client, store)
    tools.save_dictionary_draft(findings)
    with pytest.raises(ValueError, match="MAPPING_NOT_CONFIRMED"):
        tools.finalize_snapshot()
    store.update(job_id, status="awaiting_confirmation")
    preview = client.get(f"/api/v1/analyses/{job_id}/preview").json()
    stale = client.put(f"/api/v1/analyses/{job_id}/mapping", json={"version": 0, "findings": preview["findings"]})
    assert stale.status_code == 409
    assert client.post(f"/api/v1/analyses/{job_id}/confirm").status_code == 202
    assert client.post(f"/api/v1/analyses/{job_id}/confirm").status_code == 202
    store.claim(job_id)
    changed = Findings.model_validate(findings).model_dump()
    changed["tables"][0]["comment"] = "override user"
    with pytest.raises(ValueError, match="CONFIRMED_MAPPING_IMMUTABLE"):
        tools.save_findings(changed)
    snapshot_id = tools.finalize_snapshot()["snapshot_id"]
    assert tools.finalize_snapshot()["snapshot_id"] == snapshot_id
    store.update(job_id, status="succeeded")
    response = client.post(f"/api/v1/snapshots/{snapshot_id}/questions", json={"question": "有什么主键？"})
    assert response.status_code == 202
    qa_id = response.json()["task_id"]
    store.claim(qa_id)
    qa = Capabilities(store, qa_id)
    assert qa.search_metadata("users")["total"] == 1
    with pytest.raises(ValueError, match="SOURCE_TOOLS_DISABLED"):
        qa.list_objects()
    with pytest.raises(ValueError, match="UNKNOWN_EVIDENCE"):
        qa.save_answer("答案", ["fabricated"])
    evidence_id = findings["tables"][0]["evidence_ids"][0]
    qa.save_answer("users.id 为声明主键", [evidence_id])
    assert client.get("/api/v1/conversations/" + response.json()["conversation_id"]).json()["answers"]
    assert client.get(f"/api/v1/snapshots/{snapshot_id}/exports/json").status_code == 200


def test_invalid_upload_and_password_redaction(client, store):
    assert client.post("/api/v1/sources/excel", files={"file": ("bad.xlsx", b"not an xlsx")}).status_code == 400
    assert client.post("/api/v1/sources/excel", files={"file": ("script.exe", b"bad")}).status_code == 400
    secret = "test-password-not-for-model"
    body = {"name": "test", "host": "127.0.0.1", "username": "tester", "password": secret, "database": "test"}
    response = client.post("/api/v1/sources/mysql", json=body)
    assert response.status_code == 201
    assert secret not in response.text
    record = store.get(response.json()["id"], "source")
    assert secret != record["password"]
    assert store.settings.cipher.decrypt(record["password"].encode()).decode() == secret
    body["port"] = 0
    response = client.post("/api/v1/sources/mysql", json=body)
    assert response.status_code == 422 and secret not in response.text


def test_scope_budget_and_cancellation(client, store):
    job_id, tools, findings = prepared(client, store)
    other_id = store.new_job(store.get(job_id)["source_id"])
    store.claim(other_id)
    with pytest.raises(ValueError, match="UNKNOWN_EVIDENCE"):
        store.save_draft(other_id, findings)
    for _ in range(30):
        store.budget(job_id)
    with pytest.raises(ValueError, match="BUDGET_EXHAUSTED"):
        store.budget(job_id)
    client.post(f"/api/v1/analyses/{job_id}/cancel")
    with pytest.raises(ValueError, match="TASK_NOT_ACTIVE"):
        tools.inspect_workbook()
    with pytest.raises(ValueError, match="UNKNOWN_RELATIONSHIP_OBJECT"):
        Findings.model_validate(
            {
                **findings,
                "relationships": [
                    {
                        "source_table": "users",
                        "source_columns": ["id"],
                        "target_table": "ghost",
                        "target_columns": ["id"],
                        "origin": "inferred",
                        "explanation": "wrong",
                        "evidence_ids": ["x"],
                    }
                ],
            }
        )


def test_dsh_submits_lineage_and_service_only_validates_structure(client, store):
    job_id, tools, findings = prepared(client, store)
    ref = findings["tables"][0]["evidence_ids"][0]
    findings["tables"].append(
        {"name": "v_users", "kind": "view", "columns": [{"name": "user_id", "type": "BIGINT"}], "evidence_ids": [ref]}
    )
    sql_ref = store.evidence(
        job_id,
        "sql_definition",
        {"name": "v_users", "kind": "view"},
        {"sql": "WITH u AS (SELECT id FROM users) SELECT id AS user_id FROM u"},
    )
    # This edge is the DSH output. No parser or semantic analyzer runs in Python.
    findings["lineage"] = [
        {
            "source_table": "users",
            "target_table": "v_users",
            "source_column": "id",
            "target_column": "user_id",
            "expression": "u.id AS user_id",
            "evidence_ids": [sql_ref],
        }
    ]
    tools.save_findings(findings)
    assert store.get(job_id)["draft"]["lineage"] == Findings.model_validate(findings).model_dump()["lineage"]
    findings["lineage"][0]["evidence_ids"] = ["fabricated"]
    with pytest.raises(ValueError, match="UNKNOWN_EVIDENCE"):
        tools.save_findings(findings)
    findings["lineage"][0]["evidence_ids"] = [sql_ref]
    findings["lineage"][0]["source_column"] = "missing"
    with pytest.raises(ValueError, match="UNKNOWN_LINEAGE_OBJECT"):
        tools.save_findings(findings)


@pytest.mark.anyio
async def test_mcp_tool_visibility(client, store):
    from dsh.plugins.metadata_tools.server import build_server

    job_id, _, _ = prepared(client, store)
    tools = await build_server(store, job_id).list_tools()
    names = {tool.name for tool in tools}
    assert "inspect_workbook" in names and "finalize_snapshot" in names
    assert "sample_rows" not in names and "save_answer" not in names


def test_mysql_plugin_collects_evidence_without_generating_findings(store):
    source_id = store.create("source", {"type": "mysql", "config": {"database": "synthetic"}})
    job_id = store.new_job(source_id)
    store.claim(job_id)
    tools = Capabilities(store, job_id)
    fk = {"table": "orders", "constrained_columns": ["user_id"], "referred_table": "users", "referred_columns": ["id"]}
    tables = [
        {"name": "orders", "columns": [{"name": "user_id", "type": "BIGINT"}]},
        {"name": "users", "columns": [{"name": "id", "type": "BIGINT"}]},
    ]

    class Database:
        def objects(self):
            return [{"name": "orders"}, {"name": "users"}]

        def describe(self, names):
            return tables, [fk], {}, []

        def close(self):
            pass

    tools.mysql = Database
    result = tools.describe_tables(["orders", "users"])
    assert result["foreign_keys"] == [fk]
    assert store.get(job_id)["draft"]["tables"] == []
    assert store.get(job_id)["draft"]["relationships"] == []
    assert store.get(job_id)["draft"]["lineage"] == []
    ref = result["evidence_ids"][0]
    for table in tables:
        table["evidence_ids"] = [ref]
    findings = {
        "tables": tables,
        "relationships": [
            {
                "source_table": "orders",
                "source_columns": ["user_id"],
                "target_table": "users",
                "target_columns": ["id"],
                "origin": "constraint",
                "explanation": "DSH structured the raw FK evidence",
                "evidence_ids": [ref],
            }
        ],
    }
    tools.save_findings(findings)
    findings["relationships"][0]["source_table"] = "users"
    findings["relationships"][0]["source_columns"] = ["id"]
    with pytest.raises(ValueError, match="DECLARED_CONSTRAINT_WITHOUT_EVIDENCE"):
        tools.save_findings(findings)


@pytest.fixture
def anyio_backend():
    return "asyncio"

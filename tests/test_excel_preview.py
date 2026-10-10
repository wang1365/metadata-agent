import io

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def test_formatted_excel_preview_keeps_grid_merges_styles_and_clipped_anchors(client, store):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Dictionary"
    sheet.merge_cells("A1:C2")
    sheet["A1"] = "Heading"
    sheet["A1"].font = Font(name="Arial", size=14, bold=True, color="123456")
    sheet["A1"].fill = PatternFill("solid", fgColor="DDEEFF")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet["A1"].border = Border(top=Side(style="thin", color="000000"))
    sheet.column_dimensions["A"].width = 30
    sheet.row_dimensions[1].height = 30
    sheet["A3"] = 0.125
    sheet["A3"].number_format = "0.0%"
    sheet["B3"] = "=1+2"
    sheet["D4"] = "Last"
    output = io.BytesIO()
    book.save(output)
    source = client.post("/api/v1/sources/excel", files={"file": ("formatted.xlsx", output.getvalue())}).json()["id"]
    response = client.get(f"/api/v1/sources/{source}/preview?formatted=true&name=Dictionary&region=A1:D4")
    assert response.status_code == 200
    data = response.json()
    assert len(data["rows"]) == 4 and len(data["columns"]) == 4
    assert data["columns"][0]["width"] > data["columns"][1]["width"] and data["rows"][0]["height"] == 40
    cells = {(cell["row"], cell["column"]): cell for cell in data["cells"]}
    assert cells[1, 1]["rowSpan"] == 2 and cells[1, 1]["colSpan"] == 3
    assert (1, 2) not in cells and cells[2, 4]["value"] == ""
    assert cells[1, 1]["style"]["backgroundColor"] == "#ddeeff"
    assert cells[1, 1]["style"]["fontWeight"] == "bold" and cells[1, 1]["style"]["textAlign"] == "center"
    assert cells[3, 1]["value"] == "12.5%" and cells[3, 2]["value"] == "=1+2"
    clipped = client.get(f"/api/v1/sources/{source}/preview?formatted=true&region=B2:D4").json()
    assert clipped["cells"][0]["value"] == "Heading" and clipped["cells"][0]["colSpan"] == 2
    assert client.get(f"/api/v1/sources/{source}/preview?formatted=true&region=A1:ZZ999").status_code == 400
    assert store.list("job") == []
